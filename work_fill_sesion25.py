from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

src = Path(r'C:/Users/sebas/Downloads/Guia_Trabajo_Practico_Sesion_25.docx')
out = Path(r'C:/Users/sebas/OneDrive/Desktop/hola/salaMandra/Guia_Trabajo_Practico_Sesion_25_completada.docx')
doc = Document(src)

# Replace placeholder paragraph content without disturbing the supplied headings/layout.
def replace_paragraph(paragraph, text):
    for child in list(paragraph._p):
        if child.tag != qn('w:pPr'):
            paragraph._p.remove(child)
    paragraph.add_run(text)

def style_code_paragraph(p, size=8.0):
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(9.2)
    p.paragraph_format.keep_together = False
    p.paragraph_format.widow_control = False
    run = p.runs[-1]
    run.font.name = 'Consolas'
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Consolas')
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Consolas')
    run.font.size = Pt(size)

def set_cell(cell, text, size=8.5):
    p = cell.paragraphs[0]
    replace_paragraph(p, text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for r in p.runs:
        r.font.size = Pt(size)

# Team sheet: names and responsibilities from the project group.
team = doc.tables[3]
set_cell(team.cell(0,0), 'Proyecto: Sistema de alarma para habitación con Arduino UNO', 8.5)
set_cell(team.cell(0,1), 'Int. 1 (Setup/HW): Javier Ortiz', 8.5)
set_cell(team.cell(1,0), 'Int. 2 (Sensores): Sebastian Cruz', 8.5)
set_cell(team.cell(1,1), 'Int. 3 (Actuadores): Franco Arana', 8.5)
set_cell(team.cell(2,0), 'Int. 4 (Pruebas): Samuel Lima', 8.5)
set_cell(team.cell(2,1), 'Nota / pauta: simulación funcional verificada en Tinkercad; calificación pendiente del docente.', 8.0)

# Section 1 describes the implementation and provides variable declarations + setup().
replace_paragraph(doc.paragraphs[14], 'En nuestro prototipo usamos un Arduino UNO. Asignamos pines digitales al sensor PIR, al contacto magnético simulado con un interruptor, al botón de armado, al buzzer y al LED rojo. Configuramos el contacto y el botón con INPUT_PULLUP; sus entradas se activan al conectarse a GND.')
setup_code = '''const byte PIN_PIR = 2;
const byte PIN_PUERTA = 3;
const byte PIN_BOTON = 4;
const byte PIN_BUZZER = 8;
const byte PIN_LED = 9;

const unsigned long RETARDO_ENTRADA = 8000;
const unsigned long REBOTE_BOTON = 50;

bool alarmaArmada = false;
bool alarmaActiva = false;
bool entradaPendiente = false;
bool lecturaAnterior = HIGH;
bool botonEstable = HIGH;
unsigned long ultimoCambio = 0;
unsigned long inicioEntrada = 0;

void setup() {
  pinMode(PIN_PIR, INPUT);
  pinMode(PIN_PUERTA, INPUT_PULLUP);
  pinMode(PIN_BOTON, INPUT_PULLUP);
  pinMode(PIN_BUZZER, OUTPUT);
  pinMode(PIN_LED, OUTPUT);
  noTone(PIN_BUZZER);
  Serial.begin(9600);
  Serial.println("Alarma lista. Pulsa el boton para armar o desarmar.");
}'''
p = OxmlElement('w:p')
doc.paragraphs[14]._p.addnext(p)
from docx.text.paragraph import Paragraph
setup_p = Paragraph(p, doc.paragraphs[14]._parent)
run = setup_p.add_run(setup_code)
style_code_paragraph(setup_p, 8.0)

# Section 2 explains the control flow; the provided answer area receives the complete loop.
replace_paragraph(doc.paragraphs[17], 'En cada ciclo leemos el botón y los sensores. Un cambio estable del botón alterna entre alarma armada y desarmada. Si la puerta se abre, iniciamos un retardo de entrada de 8 segundos; si se detecta movimiento durante ese retardo, la alarma se activa de inmediato. Con la puerta cerrada, el movimiento también activa la alarma. Al activarse, el buzzer suena y el LED parpadea.')
loop_lines = [
'void loop() {',
'  const unsigned long ahora = millis();',
'  const bool lecturaBoton = digitalRead(PIN_BOTON);',
'',
'  // A. LECTURA DE SENSORES Y ENTRADAS',
'  if (lecturaBoton != lecturaAnterior) {',
'    ultimoCambio = ahora;',
'    lecturaAnterior = lecturaBoton;',
'  }',
'  if (ahora - ultimoCambio >= REBOTE_BOTON && lecturaBoton != botonEstable) {',
'    botonEstable = lecturaBoton;',
'    if (botonEstable == LOW) {',
'      alarmaArmada = !alarmaArmada;',
'      alarmaActiva = false;',
'      entradaPendiente = false;',
'      Serial.println(alarmaArmada ? "Alarma ARMADA" : "Alarma DESARMADA");',
'    }',
'  }',
'  // El interruptor cerrado a GND representa la puerta cerrada.',
'  const bool puertaAbierta = digitalRead(PIN_PUERTA) == HIGH;',
'  const bool movimiento = digitalRead(PIN_PIR) == HIGH;',
'',
'  // B. LOGICA DE CONTROL Y DECISIONES',
'  if (!alarmaArmada) {',
'    alarmaActiva = false;',
'    entradaPendiente = false;',
'  } else if (!alarmaActiva) {',
'    if (puertaAbierta) {',
'      if (!entradaPendiente) {',
'        entradaPendiente = true;',
'        inicioEntrada = ahora;',
'        Serial.println("Puerta abierta: inicia retardo de entrada");',
'      }',
'      if (movimiento || ahora - inicioEntrada >= RETARDO_ENTRADA) {',
'        alarmaActiva = true;',
'        Serial.println("ALARMA: intruso detectado");',
'      }',
'    } else {',
'      entradaPendiente = false;',
'      if (movimiento) {',
'        alarmaActiva = true;',
'        Serial.println("ALARMA: movimiento detectado");',
'      }',
'    }',
'  }',
'',
'  // C. ACTUACION Y ESCRITURA DE SALIDAS',
'  if (alarmaActiva) {',
'    tone(PIN_BUZZER, 2000);',
'    digitalWrite(PIN_LED, (ahora / 250) % 2);',
'  } else {',
'    noTone(PIN_BUZZER);',
'    digitalWrite(PIN_LED, alarmaArmada ? HIGH : LOW);',
'  }',
'  delay(10);',
'}'
]
code_cell = doc.tables[5].cell(0,0)
# Retain one paragraph and replace the existing blank scaffold with one code line per paragraph.
for extra in code_cell.paragraphs[1:]:
    extra._element.getparent().remove(extra._element)
first = code_cell.paragraphs[0]
replace_paragraph(first, loop_lines[0])
style_code_paragraph(first, 7.8)
for line in loop_lines[1:]:
    new_el = OxmlElement('w:p')
    first._p.addnext(new_el)
    next_p = Paragraph(new_el, first._parent)
    next_p.add_run(line if line else ' ')
    style_code_paragraph(next_p, 7.8)
    first = next_p

# Store practical document metadata without asserting a grade.
doc.core_properties.title = 'Guía de trabajo práctico sesión 25 - alarma con Arduino'
doc.core_properties.subject = 'Variables, configuración de pines y lógica de control del sistema de alarma'
doc.save(out)
print(out)
