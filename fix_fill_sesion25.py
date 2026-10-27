from docx import Document
from docx.oxml.ns import qn
p=r'C:/Users/sebas/OneDrive/Desktop/hola/salaMandra/Guia_Trabajo_Practico_Sesion_25_completada.docx'
d=Document(p)
def set_text(par, text):
    for child in list(par._p):
        if child.tag != qn('w:pPr'):
            par._p.remove(child)
    par.add_run(text)
set_text(d.paragraphs[17], '2. PROGRAMACIÓN DE LA RUTINA CÍCLICA DE CONTROL (LOOP)')
set_text(d.paragraphs[18], 'En cada ciclo leemos el botón y los sensores. Un cambio estable del botón alterna entre alarma armada y desarmada. Si la puerta se abre, iniciamos un retardo de entrada de 8 segundos; si se detecta movimiento durante ese retardo, la alarma se activa de inmediato. Con la puerta cerrada, el movimiento también activa la alarma. Al activarse, el buzzer suena y el LED parpadea.')
d.save(p)
print('corregido')
